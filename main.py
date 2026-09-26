"""
main.py - DLD Lab Word Template Generator

A single-window Tkinter (ttk) desktop app that walks the user through three steps:
    Step 1: choose how many ADDITIONAL (gate) headings the report needs (1-100). These
            come after the 3 fixed headings that every report contains.
    Step 2: type a title for each additional heading, one at a time
    Step 3: generate a .docx by copying cover_page/cover.docx and appending, in order:
              - the 3 FIXED headings (Critical Analysis, Summary of the skills learned,
                Conclusion) - each gets only a heading paragraph + 5 screenshot
                placeholder lines. NO table.
              - the user's additional ("gate") headings - each gets a heading
                paragraph + 5 screenshot placeholder lines + a 5x3 A/B/X table.
            The first heading starts on a new page (page 2) so nothing shares the
            cover page. Specific fonts/sizes are applied to fixed headings, gate
            headings, body placeholder text, and table content (see CONFIGURATION).
            After a successful save the app closes itself; on failure it stays open.

Dependencies: python-docx, tkinter (with ttk).
"""

from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

# --------------------------------------------------------------------------------------
# Imports that can fail. tkinter is required to show anything at all, so if it is
# missing we can only report on stderr. python-docx is imported defensively so that the
# GUI can display a friendly error screen instead of a raw traceback.
# --------------------------------------------------------------------------------------
try:
    import tkinter as tk
    from tkinter import filedialog, ttk
except ImportError as exc:  # pragma: no cover - depends on the Python installation
    sys.stderr.write(f"ERROR: tkinter is not available in this Python installation: {exc}\n")
    sys.exit(1)

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
except ImportError as exc:
    Document = None
    DOCX_IMPORT_ERROR = exc          # remembered so the GUI can show it
else:
    DOCX_IMPORT_ERROR = None


# ======================================================================================
# CONFIGURATION - everything you are likely to want to customise lives here.
# (See the "Customization Guide" in README.md.)
# ======================================================================================

# Folder of this script; used so the app works no matter which directory you launch from.
BASE_DIR = Path(__file__).resolve().parent

# Location of the cover-page template. Never auto-created.
COVER_PATH = BASE_DIR / "cover_page" / "cover.docx"
COVER_DISPLAY = "cover_page/cover.docx"      # short form used in messages

# Headings that appear at the start of EVERY generated document, in this exact order,
# before any heading typed by the user. Unlike user headings, these get NO table -
# only the heading paragraph and the screenshot placeholder lines.
FIXED_HEADINGS = [
    "Critical Analysis",
    "Summary of the skills learned",
    "Conclusion",
]

# Allowed range for the number of ADDITIONAL (user-defined / "gate") headings. The
# fixed headings above are not counted here.
MIN_HEADINGS = 1
MAX_HEADINGS = 100

# After a successful save, the window closes after this many milliseconds. The short
# delay lets the "Saved: <path>" status message actually be drawn before the window
# disappears. Set to 0 to close immediately.
AUTO_CLOSE_DELAY_MS = 1500

# Paragraph style used for every heading (fixed and user). Font/size overrides for each
# kind of heading are applied on top of this style - see the two blocks below.
HEADING_STYLE = "Heading 1"

# --- Formatting: the 3 fixed ("main") headings ----------------------------------------
FIXED_HEADING_FONT_NAME = "Aptos Slab Extrabold"
FIXED_HEADING_FONT_SIZE_PT = 18
FIXED_HEADING_BOLD = None          # None = no override; use the style's own weight

# --- Formatting: user-defined ("gate") headings ----------------------------------------
GATE_HEADING_FONT_NAME = "Berlin Sans FB Demi"
GATE_HEADING_FONT_SIZE_PT = 28
GATE_HEADING_BOLD = True

# --- Formatting: body text (the screenshot placeholder lines, under every heading) ----
BODY_FONT_NAME = "Calibri (Body)"
BODY_FONT_SIZE_PT = 12

# Number of empty, body-formatted paragraphs inserted under each heading for screenshots.
SCREENSHOT_PLACEHOLDER_COUNT = 5

# --- Formatting: table content (both the header row and the value rows) --------------
TABLE_FONT_NAME = "Times New Roman"
TABLE_FONT_SIZE_PT = 20

# Table style name. Must exist in cover.docx; if it does not, plain black borders are
# drawn instead. Use None for "no style, no borders". Applies only to user ("gate")
# headings - fixed headings never get a table.
TABLE_STYLE = "Table Grid"

# Table contents: first row is the header, remaining rows are the A/B presets with an
# empty X column to be filled in by the student. Table dimensions are derived from this
# list (5 rows x 3 columns by default). Every row must have the same number of cells.
TABLE_ROWS = [
    ["A", "B", "X"],
    ["0", "0", ""],
    ["0", "1", ""],
    ["1", "0", ""],
    ["1", "1", ""],
]

# --- GUI color theme (ttk "clam" base, customised) -------------------------------------
COLOR_BG = "#eef3f8"            # window / frame background
COLOR_CARD_BG = "#ffffff"       # entry / spinbox field background
COLOR_TEXT = "#1b2733"          # normal label text
COLOR_PRIMARY = "#2f6690"       # buttons, progress bar
COLOR_PRIMARY_DARK = "#1b4964"  # buttons when pressed/active
COLOR_PRIMARY_TEXT = "#ffffff"  # text on top of COLOR_PRIMARY

# Colours for inline warnings / status messages (kept distinct from the theme so
# errors/success remain readable regardless of the palette above).
ERROR_COLOR = "#b00020"
OK_COLOR = "#1b5e20"


# ======================================================================================
# DOCUMENT GENERATION (no GUI code in this section)
# ======================================================================================

class GenerationError(Exception):
    """Raised for failures whose message is already written for the end user."""


def missing_cover_message() -> str:
    """Text shown whenever cover_page/cover.docx cannot be found."""
    return (
        f"Missing required file: {COVER_DISPLAY} (expected at {COVER_PATH}). "
        "Create the 'cover_page' folder next to main.py, put 'cover.docx' inside it, "
        "and start the app again."
    )


def permission_message(path: str) -> str:
    """Text shown when Windows/OS refuses to write the output file."""
    return (
        f"Permission denied while writing '{path}'. The file may be open in Word, or "
        "the file/folder is read-only. Close the file in Word or choose another "
        "location or file name."
    )


def ensure_heading_style(doc) -> None:
    """
    Make sure the paragraph style named HEADING_STYLE exists in the document.

    Word only stores a style like "Heading 1" in a file once it has been used or
    modified, so a plain cover.docx may not contain it. If it is missing we create a
    simple built-in-style replacement (bold, 16 pt, outline level 1) so headings still
    appear in Word's Navigation Pane. The per-kind font overrides below are applied on
    top of whichever style ends up in effect, so this fallback rarely matters visually.
    """
    styles = doc.styles
    try:
        styles[HEADING_STYLE]            # KeyError if the style is not defined
        return
    except KeyError:
        pass

    style = styles.add_style(HEADING_STYLE, WD_STYLE_TYPE.PARAGRAPH, builtin=True)
    try:
        style.base_style = styles["Normal"]
    except KeyError:
        pass                              # no Normal style; leave base unset
    style.font.bold = True
    style.font.size = Pt(16)
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True

    # Outline level 0 = "Heading 1" level; appended last to respect XML element order.
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "0")
    style.element.get_or_add_pPr().append(outline)


def add_manual_borders(table) -> None:
    """
    Draw single black borders around every cell by editing the table XML directly.
    Used as a fallback when TABLE_STYLE does not exist in cover.docx.
    """
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")          # 1/8 pt units -> 0.5 pt line
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    # tblBorders must come before tblLook in the schema; insert accordingly.
    tbl_look = tbl_pr.find(qn("w:tblLook"))
    if tbl_look is not None:
        tbl_look.addprevious(borders)
    else:
        tbl_pr.append(borders)


def apply_table_style(table) -> None:
    """Apply TABLE_STYLE to the table; fall back to manual borders if it is missing."""
    if not TABLE_STYLE:
        return                                # None/"" -> leave the table unstyled
    try:
        table.style = TABLE_STYLE
    except KeyError:
        add_manual_borders(table)


def apply_font_to_run(run, font_name, size_pt, bold=None) -> None:
    """Set name/size/bold on a single run. Each argument is skipped when falsy/None."""
    if font_name:
        run.font.name = font_name
    if size_pt:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold


def apply_font(paragraph, font_name, size_pt, bold=None) -> None:
    """Apply the same name/size/bold override to every run already in a paragraph."""
    for run in paragraph.runs:
        apply_font_to_run(run, font_name, size_pt, bold)


def add_heading_block(doc, title: str, is_fixed: bool, start_on_new_page: bool = False) -> None:
    """
    Append one complete section to the document.

    Fixed headings (is_fixed=True): heading paragraph + screenshot placeholder lines.
        NO table.
    User / "gate" headings (is_fixed=False): heading paragraph + screenshot placeholder
        lines + a 5x3 A/B/X table.

    start_on_new_page: when True the heading paragraph gets Word's "Page break before"
    property. Used for the very first heading so it always begins on page 2, with
    nothing (not even an empty paragraph) added to the cover page itself.
    """
    # 1. Heading paragraph in the configured style, with the font that matches its kind.
    heading = doc.add_paragraph(title, style=HEADING_STYLE)
    if is_fixed:
        apply_font(heading, FIXED_HEADING_FONT_NAME, FIXED_HEADING_FONT_SIZE_PT, FIXED_HEADING_BOLD)
    else:
        apply_font(heading, GATE_HEADING_FONT_NAME, GATE_HEADING_FONT_SIZE_PT, GATE_HEADING_BOLD)
    if start_on_new_page:
        heading.paragraph_format.page_break_before = True

    # 2. Empty, body-formatted lines where the student pastes screenshots. An empty run
    # is added explicitly (add_paragraph() with no text creates no run at all) so the
    # body font is actually attached to the line.
    for _ in range(SCREENSHOT_PLACEHOLDER_COUNT):
        placeholder = doc.add_paragraph()
        run = placeholder.add_run("")
        apply_font_to_run(run, BODY_FONT_NAME, BODY_FONT_SIZE_PT)

    # 3. Truth table - user ("gate") headings only.
    if not is_fixed:
        table = doc.add_table(rows=len(TABLE_ROWS), cols=len(TABLE_ROWS[0]))
        apply_table_style(table)
        for r_index, row_values in enumerate(TABLE_ROWS):
            for c_index, value in enumerate(row_values):
                cell = table.cell(r_index, c_index)
                cell.text = value
                for paragraph in cell.paragraphs:
                    apply_font(paragraph, TABLE_FONT_NAME, TABLE_FONT_SIZE_PT)


def build_document_content(doc, user_titles: list[str]) -> None:
    """
    Validate the table configuration, then append one block per heading, in order:
    first the FIXED_HEADINGS (no table), then the user's additional "gate" headings
    (each with a table). The first block starts on a new page so the cover page stays
    clean.
    """
    columns = len(TABLE_ROWS[0]) if TABLE_ROWS else 0
    if columns == 0 or any(len(row) != columns for row in TABLE_ROWS):
        raise GenerationError(
            "TABLE_ROWS in main.py is misconfigured: it must be non-empty and every "
            "row must have the same number of cells."
        )
    ensure_heading_style(doc)

    fixed_count = len(FIXED_HEADINGS)
    all_titles = list(FIXED_HEADINGS) + list(user_titles)
    for position, title in enumerate(all_titles):
        add_heading_block(
            doc, title,
            is_fixed=(position < fixed_count),
            start_on_new_page=(position == 0),
        )


def _remove_quietly(path: str) -> None:
    """Best-effort delete of a partially created output file; never raises."""
    try:
        Path(path).unlink()
    except OSError:
        pass


def generate_document(out_path: str, user_titles: list[str]) -> None:
    """
    Create the report template at out_path.

    user_titles are the ADDITIONAL ("gate") headings typed in Step 2; the fixed
    headings are added automatically in front of them.

    Steps: check the cover exists -> shutil.copy it to out_path -> open the copy with
    python-docx -> append the heading blocks -> save.
    Raises GenerationError (with a user-ready message) for every anticipated failure.
    Any unexpected exception is allowed to propagate to the caller.
    """
    # Re-check here: the file could have been deleted since start-up.
    if not COVER_PATH.is_file():
        raise GenerationError(missing_cover_message())

    # --- Copy the cover page to the chosen destination ---------------------------------
    # NOTE: SameFileError and PermissionError are subclasses of OSError, so they must be
    # caught before the generic OSError handler.
    try:
        shutil.copy(COVER_PATH, out_path)
    except shutil.SameFileError:
        raise GenerationError(
            "Cannot save over the template itself. Choose a different file name or folder."
        ) from None
    except PermissionError:
        raise GenerationError(permission_message(out_path)) from None
    except OSError as exc:
        raise GenerationError(f"Could not write '{out_path}': {exc.strerror or exc}") from exc

    # --- From here on a file exists at out_path; if anything fails we delete it so the
    # --- user is not left with a half-made document.
    try:
        # Open the copy. Any failure other than a permission problem means the
        # cover file is not a readable .docx.
        try:
            doc = Document(out_path)
        except PermissionError:
            raise GenerationError(permission_message(out_path)) from None
        except Exception as exc:
            raise GenerationError(
                f"{COVER_DISPLAY} could not be read - it may be corrupt or not a real "
                f".docx file ({type(exc).__name__}: {exc})."
            ) from exc

        # Append the per-heading content (fixed headings first, then the user's).
        build_document_content(doc, user_titles)

        # Save in place.
        try:
            doc.save(out_path)
        except PermissionError:
            raise GenerationError(permission_message(out_path)) from None
        except OSError as exc:
            raise GenerationError(f"Could not save '{out_path}': {exc.strerror or exc}") from exc
    except Exception:
        _remove_quietly(out_path)
        raise


# ======================================================================================
# INPUT VALIDATION (pure functions, no GUI)
# ======================================================================================

def parse_heading_count(text: str):
    """
    Validate the Step 1 input (number of additional headings).
    Returns (value, None) when valid, or (None, warning_message) when invalid.
    Only plain ASCII digits are accepted: no signs, decimals, spaces inside, or letters.
    """
    cleaned = text.strip()
    if not cleaned:
        return None, f"Enter a whole number from {MIN_HEADINGS} to {MAX_HEADINGS}."
    if not re.fullmatch(r"[0-9]+", cleaned):
        return None, "Digits only - no letters, decimals or signs."
    if len(cleaned) > 9:                       # avoid absurdly long numbers
        return None, f"Number must be between {MIN_HEADINGS} and {MAX_HEADINGS}."
    value = int(cleaned)
    if value < MIN_HEADINGS or value > MAX_HEADINGS:
        return None, f"Number must be between {MIN_HEADINGS} and {MAX_HEADINGS}."
    return value, None


# ======================================================================================
# GUI
# ======================================================================================

TOTAL_STEPS = 3  # used only for the "Step X of 3" progress label


class App:
    """
    Single-window wizard. The top area ("content") is cleared and rebuilt for each step;
    a persistent progress label sits above it and a persistent status label sits below.
    """

    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.heading_count = None          # int once Step 1 has been passed (user headings only)
        self.titles: list[str] = []        # one entry per user heading ("" = not entered yet)
        self.current_index = 0             # 0-based user heading being edited in Step 2
        self.fatal = False                 # True when a start-up error stopped the app
        self.step3_buttons: list = []      # Step 3 buttons, disabled while auto-closing

        # Window setup: resizable, with a sensible minimum size.
        root.title("DLD Lab Word Template Generator")
        root.geometry("560x300")
        root.minsize(440, 260)
        root.columnconfigure(0, weight=1)
        root.rowconfigure(1, weight=1)     # content area grows; progress/separator/status don't
        root.configure(background=COLOR_BG)

        self._apply_theme()

        # Persistent progress indicator ("Step X of 3"), above the content area.
        self.progress_var = tk.StringVar(value="")
        self.progress_label = ttk.Label(
            root, textvariable=self.progress_var, style="Progress.TLabel", anchor="w"
        )
        self.progress_label.grid(row=0, column=0, sticky="ew")

        # Content area (swapped per step).
        self.content = ttk.Frame(root, padding=16)
        self.content.grid(row=1, column=0, sticky="nsew")
        self.content.columnconfigure(0, weight=1)
        self.content.rowconfigure(9, weight=1)   # spacer row pushes buttons (row 10) down

        # Separator + single status label at the bottom of the window.
        ttk.Separator(root, orient="horizontal").grid(row=2, column=0, sticky="ew")
        self.status_label = ttk.Label(root, text="", anchor="w", justify="left", padding=(16, 6))
        self.status_label.grid(row=3, column=0, sticky="ew")
        # Keep the status text wrapped to the current window width (long paths/errors).
        self.status_label.bind(
            "<Configure>",
            lambda event: self.status_label.configure(wraplength=max(event.width - 32, 100)),
        )

        # Start-up checks: fail early with a clear message instead of a traceback.
        problem = self.startup_problem()
        if problem:
            self.show_fatal(problem)
        else:
            self.show_step1()

    # ---------------------------------------------------------------- theming ---------

    def _apply_theme(self) -> None:
        """
        Set up a single, consistent color palette (a soft blue/steel theme) across every
        ttk widget in the app. Uses the built-in 'clam' theme as a base because it is the
        most reliable one to recolor consistently across Windows/macOS/Linux.
        """
        style = ttk.Style(self.root)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass  # 'clam' should always be available; if not, keep the platform default

        style.configure(".", background=COLOR_BG, foreground=COLOR_TEXT)
        style.configure("TFrame", background=COLOR_BG)
        style.configure("TLabel", background=COLOR_BG, foreground=COLOR_TEXT)
        style.configure("TSeparator", background=COLOR_PRIMARY)

        style.configure(
            "TButton", background=COLOR_PRIMARY, foreground=COLOR_PRIMARY_TEXT,
            padding=(12, 6), borderwidth=0, focusthickness=3, focuscolor=COLOR_PRIMARY_DARK,
        )
        style.map(
            "TButton",
            background=[("active", COLOR_PRIMARY_DARK), ("disabled", "#a9b7c4")],
            foreground=[("disabled", "#e7edf3")],
        )

        style.configure("TEntry", fieldbackground=COLOR_CARD_BG, foreground=COLOR_TEXT)
        style.configure("TSpinbox", fieldbackground=COLOR_CARD_BG, foreground=COLOR_TEXT)

        # The progress bar/label at the very top of the window: solid accent background.
        style.configure(
            "Progress.TLabel", background=COLOR_PRIMARY, foreground=COLOR_PRIMARY_TEXT,
            padding=(16, 8), font=("TkDefaultFont", 10, "bold"),
        )

    # ---------------------------------------------------------------- helpers ---------

    def startup_problem(self):
        """Return an error message if the app cannot run, else None."""
        if DOCX_IMPORT_ERROR is not None:
            return (
                "The 'python-docx' package is not installed (or failed to import). "
                "Run:  pip install python-docx   "
                "(install 'python-docx', not 'docx').  "
                f"Details: {DOCX_IMPORT_ERROR}"
            )
        if not COVER_PATH.is_file():
            return missing_cover_message()
        return None

    def set_status(self, text: str, kind: str = "") -> None:
        """Update the bottom status label. kind: 'error', 'ok' or '' (neutral)."""
        color = {"error": ERROR_COLOR, "ok": OK_COLOR}.get(kind, COLOR_TEXT)
        self.status_label.configure(text=text, foreground=color)

    def _reset_content(self) -> None:
        """Destroy the widgets of the current step and clear the status line."""
        for child in self.content.winfo_children():
            child.destroy()
        self.step3_buttons = []
        self.set_status("")

    def _add_wrapping_label(self, text: str, row: int, foreground: str = "") -> ttk.Label:
        """
        Create a left-aligned label whose text re-wraps to the width of the window.
        Used for longer texts (Step 1 prompt, fatal error message).
        """
        options = {"text": text, "justify": "left"}
        if foreground:
            options["foreground"] = foreground
        label = ttk.Label(self.content, **options)
        label.grid(row=row, column=0, sticky="ew")
        label.bind(
            "<Configure>",
            lambda event: label.configure(wraplength=max(event.width - 10, 100)),
        )
        return label

    def _add_warning_label(self, row: int) -> tk.StringVar:
        """Create an (initially empty) red inline-warning label and return its variable."""
        var = tk.StringVar(value="")
        label = ttk.Label(self.content, textvariable=var, foreground=ERROR_COLOR, justify="left")
        label.grid(row=row, column=0, sticky="w", pady=(4, 0))
        return var

    def _add_buttons(self, specs):
        """Place right-aligned buttons at the bottom of the content area. specs = [(text, command)]."""
        frame = ttk.Frame(self.content)
        frame.grid(row=10, column=0, sticky="e", pady=(12, 0))
        buttons = []
        for column, (text, command) in enumerate(specs):
            button = ttk.Button(frame, text=text, command=command)
            button.grid(row=0, column=column, padx=(6, 0))
            buttons.append(button)
        return buttons

    # ------------------------------------------------------------- fatal screen -------

    def show_fatal(self, message: str) -> None:
        """
        Show an unrecoverable start-up error inside the main window (no popup) with an
        Exit button. The message is also written to stderr, and the process exits with
        status 1 once the window is closed.
        """
        self.fatal = True
        sys.stderr.write(f"ERROR: {message}\n")
        self.progress_var.set("Setup Error")
        self._reset_content()
        self._add_wrapping_label(message, row=0, foreground=ERROR_COLOR)
        exit_button, = self._add_buttons([("Exit", self.root.destroy)])
        exit_button.focus_set()
        self.set_status("Fatal error - the application cannot continue. Click Exit to close.", "error")

    # ------------------------------------------------------------------ Step 1 -------

    def show_step1(self) -> None:
        """Step 1: ask for the number of ADDITIONAL headings (integer 1-100)."""
        self.progress_var.set(f"Step 1 of {TOTAL_STEPS}")
        self._reset_content()
        self._add_wrapping_label(
            f"Number of additional headings (after the {len(FIXED_HEADINGS)} fixed ones):",
            row=0,
        )

        # Pre-fill with the stored count when returning from Step 2 ("preserve if valid").
        initial = str(self.heading_count) if self.heading_count else str(MIN_HEADINGS)
        self.count_var = tk.StringVar(value=initial)

        # Warning label is created BEFORE the spinbox because the key-validation
        # callbacks below write to it.
        self.count_warning = self._add_warning_label(row=2)

        # Key-level validation: only up to 3 ASCII digits can be typed/pasted.
        vcmd = (self.root.register(self._on_count_key), "%P")
        icmd = (self.root.register(self._on_count_key_rejected),)
        self.count_spin = ttk.Spinbox(
            self.content, from_=MIN_HEADINGS, to=MAX_HEADINGS, increment=1, width=8,
            textvariable=self.count_var,
            validate="key", validatecommand=vcmd, invalidcommand=icmd,
        )
        self.count_spin.grid(row=1, column=0, sticky="w", pady=(6, 0))
        self.count_spin.bind("<Return>", lambda event: self.on_step1_next())
        self.count_spin.focus_set()

        self._add_buttons([("Next", self.on_step1_next)])    # no Back button on Step 1

    def _on_count_key(self, proposed: str) -> bool:
        """validatecommand: allow the edit only if the result is empty or 1-3 ASCII digits."""
        allowed = re.fullmatch(r"[0-9]{0,3}", proposed) is not None
        if allowed:
            self.count_warning.set("")
        return allowed

    def _on_count_key_rejected(self) -> None:
        """invalidcommand: tell the user why the keystroke was ignored."""
        self.count_warning.set("Only whole numbers are allowed (digits 0-9, up to 3 digits).")

    def on_step1_next(self) -> None:
        """Validate the count; advance to Step 2 only if valid."""
        value, warning = parse_heading_count(self.count_var.get())
        if warning:
            self.count_warning.set(warning)
            return
        self.count_warning.set("")

        # Resize the title list to the new count, keeping titles already typed.
        self.titles = self.titles[:value] + [""] * (value - len(self.titles))
        self.heading_count = value
        self.current_index = 0
        self.show_step2()

    # ------------------------------------------------------------------ Step 2 -------

    def show_step2(self) -> None:
        """Step 2: ask for the title of user heading number current_index+1 (of N)."""
        self.progress_var.set(f"Step 2 of {TOTAL_STEPS}")
        self._reset_content()
        total = self.heading_count
        index = self.current_index
        ttk.Label(
            self.content, text=f"Heading {index + 1} of {total} \u2014 enter title:"
        ).grid(row=0, column=0, sticky="w")

        # Pre-fill with any previously entered title for this heading.
        self.title_var = tk.StringVar(value=self.titles[index])
        self.title_entry = ttk.Entry(self.content, textvariable=self.title_var)
        self.title_entry.grid(row=1, column=0, sticky="ew", pady=(6, 0))
        self.title_entry.bind("<Return>", lambda event: self.on_step2_next())
        self.title_entry.focus_set()
        self.title_entry.icursor("end")

        self.title_warning = self._add_warning_label(row=2)
        self._add_buttons([("Back", self.on_step2_back), ("Next", self.on_step2_next)])

    def on_step2_back(self) -> None:
        """Return to Step 1. A non-blank title being typed is kept for later."""
        typed = self.title_var.get().strip()
        if typed:
            self.titles[self.current_index] = typed
        self.show_step1()

    def on_step2_next(self) -> None:
        """Validate the title (non-empty), store it, and move on (Step 3 after the last one)."""
        title = self.title_var.get().strip()
        if not title:
            self.title_warning.set("Title cannot be empty.")
            return
        self.titles[self.current_index] = title

        if self.current_index < self.heading_count - 1:
            self.current_index += 1
            self.show_step2()
        else:
            self.show_step3()

    # ------------------------------------------------------------------ Step 3 -------

    def show_step3(self) -> None:
        """Step 3: summary (additional / fixed / total headings) + Generate button."""
        self.progress_var.set(f"Step 3 of {TOTAL_STEPS}")
        self._reset_content()
        fixed = len(FIXED_HEADINGS)
        summary = (
            f"Number of additional headings: {self.heading_count}\n"
            f"Fixed headings: {fixed}\n"
            f"Total headings in the document: {fixed + self.heading_count}"
        )
        ttk.Label(self.content, text=summary, justify="left").grid(row=0, column=0, sticky="w")

        back_button, generate_button = self._add_buttons(
            [("Back", self.on_step3_back), ("Generate", self.on_generate)]
        )
        self.step3_buttons = [back_button, generate_button]
        generate_button.focus_set()

    def on_step3_back(self) -> None:
        """Return to Step 2 at the last user heading."""
        self.current_index = self.heading_count - 1
        self.show_step2()

    def schedule_auto_close(self) -> None:
        """
        Close the application after a successful save. The buttons are disabled first
        so nothing can be clicked while the (short) delay runs; the delay itself lets
        the success message be drawn before the window goes away.
        """
        for button in self.step3_buttons:
            button.state(["disabled"])
        self.root.update_idletasks()               # make sure the status text is painted
        if AUTO_CLOSE_DELAY_MS <= 0:
            self.root.destroy()
        else:
            self.root.after(AUTO_CLOSE_DELAY_MS, self.root.destroy)

    def on_generate(self) -> None:
        """
        Ask where to save, build the document, and report the result in the status label.
        On success the app then closes itself; on failure the window stays open.
        """
        out_path = filedialog.asksaveasfilename(
            parent=self.root,
            title="Save Word document",
            defaultextension=".docx",
            filetypes=[("Word document", "*.docx")],
        )
        # Cancelling returns "" (or an empty tuple on some platforms).
        if not out_path:
            self.set_status("Save cancelled - no file was created.")
            return

        # Some platforms do not append the default extension; make sure it is .docx.
        if not out_path.lower().endswith(".docx"):
            out_path += ".docx"

        try:
            generate_document(out_path, self.titles)
        except GenerationError as exc:
            self.set_status(str(exc), "error")     # stay open so the error can be read
        except Exception as exc:                   # anything we did not anticipate
            self.set_status(f"Unexpected error ({type(exc).__name__}): {exc}", "error")
        else:
            self.set_status(f"Saved: {out_path}", "ok")
            self.schedule_auto_close()


# ======================================================================================
# ENTRY POINT
# ======================================================================================

def main() -> int:
    """Create the window, run the event loop, and return the process exit code."""
    root = tk.Tk()
    app = App(root)
    root.mainloop()
    return 1 if app.fatal else 0


if __name__ == "__main__":
    sys.exit(main())